"""Patients from an Excel or CSV file.

What these pin: a CSV saved by Excel on Windows (cp1252, ``;``) and an
``.xlsx`` read the same; columns are recognised by their usual names,
each field once; dates follow the order the person chose, phones come out
in international format, sex from the words people use, and a value that
does not fit is dropped and named in the file's warnings; the template
opens with the columns the import recognises. Over HTTP the sheet becomes
a ``.dpm`` job that validates and imports like any other, and importing
the same file again creates nobody twice.
"""

import io
import json
from pathlib import Path
from uuid import uuid4

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth.models import Clinic
from app.modules.migration_import import service as import_service
from app.modules.migration_import.dpmf import open_dpmf
from app.modules.migration_import.models import ImportJob
from app.modules.migration_import.sheets import (
    SheetError,
    convert_patients,
    read_sheet,
    suggest_fields,
    template_bytes,
)
from app.modules.patients.models import Patient

SHEETS = "/api/v1/migration_import/sheets"
JOBS = "/api/v1/migration_import/jobs"

HEADERS = [
    "Nombre(s)",
    "Apellido Paterno",
    "Apellido Materno",
    "Fecha Nac.",
    "Sexo",
    "Celular",
    "Correo",
    "Edad",
]
ROWS = [
    ["José", "Núñez", "Peña", "03/04/1985", "H", "55 1234 5678", "JOSE@Example.com", "40"],
    ["Ana", "López", "", "31/12/1990", "Mujer", "+34 612 345 678", "no-es-correo", "35"],
    ["", "", "", "", "", "", "", ""],
    ["", "Sin nombre", "", "", "", "", "", ""],
    ["Luis", "Ruiz", "", "1990-13-45", "?", "123", "", ""],
]


def _csv(path: Path) -> Path:
    lines = [";".join(HEADERS)] + [";".join(row) for row in ROWS]
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("cp1252"))
    return path


def _xlsx(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    book.save(path)
    return path


def test_csv_and_xlsx_read_the_same(tmp_path: Path) -> None:
    from_csv = read_sheet(_csv(tmp_path / "pacientes.csv"))
    from_xlsx = read_sheet(_xlsx(tmp_path / "pacientes.xlsx"))
    assert from_csv.headers == from_xlsx.headers == HEADERS
    assert from_csv.rows == from_xlsx.rows
    assert from_csv.rows[0][0] == "José"  # cp1252 decoded
    assert from_csv.lines == [2, 3, 5, 6]  # the empty row is skipped, lines kept

    (tmp_path / "x.pdf").write_bytes(b"%PDF")
    with pytest.raises(SheetError):
        read_sheet(tmp_path / "x.pdf")
    (tmp_path / "solo.csv").write_text("Nombre\n")
    with pytest.raises(SheetError):
        read_sheet(tmp_path / "solo.csv")


def test_columns_are_recognised_by_their_usual_names() -> None:
    assert suggest_fields(HEADERS) == [
        "given_name",
        "family_name",
        "second_family_name",
        "date_of_birth",
        "sex",
        "mobile_phone",
        "email",
        None,  # "Edad" is not a field
    ]
    # One field, one column: the best match keeps it.
    assert suggest_fields(["Nombre", "Nombre(s)"]) == ["given_name", None]


def test_values_are_read_as_the_person_said(tmp_path: Path) -> None:
    sheet = read_sheet(_csv(tmp_path / "pacientes.csv"))
    mapping = dict(enumerate(suggest_fields(sheet.headers)))
    mapping = {i: f for i, f in mapping.items() if f}
    target = tmp_path / "out.dpm"
    counts = convert_patients(
        sheet, mapping, date_order="dmy", region="MX", digest="abc", destination=target, label="x"
    )
    assert counts == {"patient": 3}

    with open_dpmf(target) as handle:
        assert handle.recompute_integrity_hash() == handle.declared_integrity_hash()
        payloads = {
            json.loads(r[3])["given_name"]: json.loads(r[3]) for r in handle.entity_iter("patient")
        }
        warnings = {(w[1], w[3]) for w in handle.warnings_iter()}

    jose = payloads["José"]
    assert jose["family_name"] == "Núñez Peña"
    assert jose["date_of_birth"] == "1985-04-03"  # 3 April: day first, as chosen
    assert jose["sex"] == "male"
    assert jose["mobile_phone"] == "+525512345678"
    assert jose["email"] == "jose@example.com"
    ana = payloads["Ana"]
    assert ana["mobile_phone"] == "+34612345678"  # its own prefix wins
    assert "email" not in ana
    luis = payloads["Luis"]
    assert {"date_of_birth", "sex", "mobile_phone"}.isdisjoint(luis)
    assert ("fila 3", "sheet.invalid_email") in warnings
    assert ("fila 5", "sheet.no_name") in warnings
    assert {("fila 6", "sheet.invalid_date_of_birth"), ("fila 6", "sheet.invalid_sex")} <= warnings

    # Month first, and the same cell is another day.
    counts = convert_patients(
        sheet,
        mapping,
        date_order="mdy",
        region="MX",
        digest="abc",
        destination=tmp_path / "mdy.dpm",
        label="x",
    )
    with open_dpmf(tmp_path / "mdy.dpm") as handle:
        born = {
            json.loads(r[3])["given_name"]: json.loads(r[3]).get("date_of_birth")
            for r in handle.entity_iter("patient")
        }
    assert born["José"] == "1985-03-04"


def test_the_template_has_the_columns_the_import_reads() -> None:
    book = load_workbook(io.BytesIO(template_bytes()))
    headers = [cell.value for cell in book["Pacientes"][1]]
    assert headers[:2] == ["Nombre", "Apellidos"]
    assert None not in suggest_fields(headers)  # every template column is recognised
    assert "Instrucciones" in book.sheetnames


@pytest.mark.asyncio
async def test_a_sheet_becomes_a_job_and_imports_once(
    client: AsyncClient,
    auth_headers: dict,
    db_session: AsyncSession,
    test_clinic: Clinic,
    tmp_path: Path,
) -> None:
    template = await client.get(f"{SHEETS}/template", headers=auth_headers)
    assert template.status_code == 200
    assert template.content[:2] == b"PK"

    refused = await client.post(
        SHEETS, headers=auth_headers, files={"file": ("x.pdf", b"%PDF", "application/pdf")}
    )
    assert refused.status_code == 422

    data = _xlsx(tmp_path / "pacientes.xlsx").read_bytes()

    async def upload_and_convert() -> str:
        up = await client.post(
            SHEETS,
            headers=auth_headers,
            files={"file": ("pacientes.xlsx", data, "application/octet-stream")},
        )
        assert up.status_code == 201, up.text
        body = up.json()["data"]
        assert body["row_count"] == 4 and body["job"]["status"] == "mapping"
        mapping = {
            str(c["index"]): c["suggested_field"] for c in body["columns"] if c["suggested_field"]
        }
        done = await client.post(
            f"{SHEETS}/{body['job']['id']}/convert",
            headers=auth_headers,
            json={"mapping": mapping, "date_order": "dmy"},
        )
        assert done.status_code == 200, done.text
        assert done.json()["data"]["status"] == "uploaded"
        job_id = body["job"]["id"]
        valid = await client.post(f"{JOBS}/{job_id}/validate", headers=auth_headers, json={})
        assert valid.json()["data"]["status"] == "validated", valid.text
        return job_id

    async def run(job_id: str) -> None:
        job = (
            await db_session.execute(select(ImportJob).where(ImportJob.id == job_id))
        ).scalar_one()
        await import_service.ImportJobService._run_pipeline(db_session, job, passphrase=None)

    async def patients() -> int:
        return (
            await db_session.execute(
                select(func.count(Patient.id)).where(Patient.clinic_id == test_clinic.id)
            )
        ).scalar_one()

    before = await patients()
    await run(await upload_and_convert())
    assert await patients() == before + 3
    # The same file again: same row ids, nobody created twice.
    await run(await upload_and_convert())
    assert await patients() == before + 3

    # A mapping without a name is refused before anything is written.
    up = await client.post(
        SHEETS, headers=auth_headers, files={"file": ("p.xlsx", data, "application/octet-stream")}
    )
    job_id = up.json()["data"]["job"]["id"]
    nameless = await client.post(
        f"{SHEETS}/{job_id}/convert", headers=auth_headers, json={"mapping": {"6": "email"}}
    )
    assert nameless.status_code == 422
    twice = await client.post(
        f"{SHEETS}/{job_id}/convert",
        headers=auth_headers,
        json={"mapping": {"0": "given_name", "1": "given_name"}},
    )
    assert twice.status_code == 422


def test_a_column_is_recognised_by_what_it_holds() -> None:
    """A header that says nothing: the values do. The name still wins."""
    from app.modules.migration_import.sheets.common import suggest, suggest_by_content
    from app.modules.migration_import.sheets.kinds import KINDS

    headers = ["Nombre", "Dato 1", "Dato 2", "Dato 3", "Dato 4"]
    rows = [
        ["Ana", "ana@example.com", "NUPJ850403HDFXXX01", "25/03/1980", "55 1234 5678"],
        ["Luis", "luis@example.com", "LOPA900101MDFXXX02", "02/11/1975", "55 8765 4321"],
    ]
    spec = KINDS["patients"]
    by_name = suggest(headers, spec.fields)
    assert by_name == ["given_name", None, None, None, None]
    assert suggest_by_content(rows, by_name, spec.by_kind, "MX") == [
        "given_name",
        "email",
        "national_id",
        "date_of_birth",
        "mobile_phone",
    ]


def test_the_order_of_dates_is_read_when_the_values_give_it_away() -> None:
    from app.modules.migration_import.sheets.common import infer_date_order

    assert infer_date_order(["03/04/1985", "25/03/1990"]) == "dmy"
    assert infer_date_order(["03/04/1985", "03/25/1990"]) == "mdy"
    assert infer_date_order(["1985/04/03"]) == "ymd"
    # All fit both ways: the person chooses.
    assert infer_date_order(["03/04/1985", "01/02/1990"]) is None
    # Excel date cells read the same either way and decide nothing.
    assert infer_date_order(["1985-04-03", "25/03/1990"]) == "dmy"
    # Contradictory: nothing is assumed.
    assert infer_date_order(["25/03/1990", "03/25/1990"]) is None


def test_a_sheet_without_headers_is_read_from_its_first_row(tmp_path: Path) -> None:
    path = tmp_path / "sin-encabezados.csv"
    path.write_text("Ana;López;ana@example.com;25/03/1980\nLuis;Ruiz;luis@example.com;02/11/1975\n")
    sheet = read_sheet(path)
    assert sheet.headerless
    assert sheet.headers == ["Columna 1", "Columna 2", "Columna 3", "Columna 4"]
    assert [row[0] for row in sheet.rows] == ["Ana", "Luis"]
    assert sheet.lines == [1, 2]
    # Names in the first row are names.
    assert not read_sheet(_csv(tmp_path / "con.csv")).headerless


@pytest.mark.asyncio
async def test_what_goes_wrong_is_said_in_spanish_and_handed_back_as_a_sheet(
    client: AsyncClient, auth_headers: dict, test_clinic: Clinic, tmp_path: Path
) -> None:
    empty = await client.post(
        SHEETS, headers=auth_headers, files={"file": ("vacio.csv", b"", "text/csv")}
    )
    assert empty.status_code == 422
    assert empty.json()["message"] == "El archivo está vacío."

    data = _xlsx(tmp_path / "pacientes.xlsx").read_bytes()
    up = await client.post(
        SHEETS, headers=auth_headers, files={"file": ("p.xlsx", data, "application/octet-stream")}
    )
    body = up.json()["data"]
    assert body["date_order"] == "dmy"  # "31/12/1990" gave it away
    assert body["headerless"] is False
    job_id = body["job"]["id"]
    nameless = await client.post(
        f"{SHEETS}/{job_id}/convert", headers=auth_headers, json={"mapping": {"6": "email"}}
    )
    assert nameless.status_code == 422
    assert "Nombre" in nameless.json()["message"]

    mapping = {
        str(c["index"]): c["suggested_field"] for c in body["columns"] if c["suggested_field"]
    }
    done = await client.post(
        f"{SHEETS}/{job_id}/convert",
        headers=auth_headers,
        json={"mapping": mapping, "date_order": "dmy"},
    )
    assert done.status_code == 200, done.text
    report = await client.get(f"{SHEETS}/{job_id}/problems", headers=auth_headers)
    assert report.status_code == 200
    book = load_workbook(io.BytesIO(report.content))
    rows = list(book["Filas con problemas"].iter_rows(values_only=True))
    assert rows[0][:2] == ("Fila", "Motivo")
    lines = {row[0]: row for row in rows[1:]}
    assert set(lines) == {3, 5, 6}  # bad email; no name; bad date and sex
    assert "Correo no válido" in lines[3][1]
    assert lines[3][2] == "Ana"  # the row as it was, to fix
    assert "Instrucciones" in book.sheetnames


@pytest.mark.asyncio
async def test_a_patient_the_clinic_already_has_is_not_created_twice(
    client: AsyncClient,
    auth_headers: dict,
    db_session: AsyncSession,
    test_clinic: Clinic,
) -> None:
    """Sure duplicates (ID document; name with date of birth, email or phone)
    are left out by default and imported only when asked; a shared name
    alone is imported and flagged; a person twice in the file comes once."""
    from datetime import date as day

    clinic_id = test_clinic.id
    db_session.add_all(
        [
            Patient(
                id=uuid4(),
                clinic_id=clinic_id,
                first_name="José",
                last_name="Núñez",
                national_id="NUPJ850403HDFXXX01",
            ),
            Patient(
                id=uuid4(),
                clinic_id=clinic_id,
                first_name="Ana",
                last_name="Ruiz",
                date_of_birth=day(1990, 12, 31),
            ),
            Patient(id=uuid4(), clinic_id=clinic_id, first_name="Luis", last_name="Gómez"),
        ]
    )
    await db_session.commit()
    csv = "\n".join(
        [
            "Nombre;Apellidos;CURP;Fecha de nacimiento;Correo",
            "José;Núñez Peña;NUPJ850403HDFXXX01;;",  # same document
            "Ana;Ruiz;;31/12/1990;",  # same name and date of birth
            "Luis;Gómez;;;",  # same name only
            "Marta;Sol;;;marta@example.com",
            "Marta;Sol;;;marta@example.com",  # twice in the file
        ]
    )

    async def convert(on_duplicate: str | None) -> tuple[set[tuple[str, str]], str]:
        up = await client.post(
            SHEETS,
            headers=auth_headers,
            files={"file": ("p.csv", csv.encode("utf-8"), "text/csv")},
        )
        body = up.json()["data"]
        mapping = {
            str(c["index"]): c["suggested_field"] for c in body["columns"] if c["suggested_field"]
        }
        extra = {"on_duplicate": on_duplicate} if on_duplicate else {}
        done = await client.post(
            f"{SHEETS}/{body['job']['id']}/convert",
            headers=auth_headers,
            json={"mapping": mapping, "date_order": "dmy", **extra},
        )
        assert done.status_code == 200, done.text
        job = (
            await db_session.execute(select(ImportJob).where(ImportJob.id == body["job"]["id"]))
        ).scalar_one()
        with open_dpmf(Path(job.file_path)) as handle:
            warnings = {(w[1], w[3]) for w in handle.warnings_iter()}
            names = {json.loads(r[3])["given_name"] for r in handle.entity_iter("patient")}
        return warnings, ",".join(sorted(names))

    warnings, names = await convert(None)
    assert names == "Luis,Marta"  # José and Ana are already there; Marta once
    assert ("fila 2", "sheet.duplicate") in warnings
    assert ("fila 3", "sheet.duplicate") in warnings
    assert ("fila 4", "sheet.possible_duplicate") in warnings
    assert ("fila 6", "sheet.repeated_in_file") in warnings

    warnings, names = await convert("import")
    assert names == "Ana,José,Luis,Marta"
    assert ("fila 2", "sheet.duplicate_imported") in warnings
