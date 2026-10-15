"""The rows of a sheet that did not go in whole, as a sheet to fix.

A copy of the uploaded columns with only the rows that had a problem, the
line each was on and why. The person fixes them there and uploads that
file: the rows that went in already are not created twice.
"""

from __future__ import annotations

import io
import re

from .reader import Sheet

_LINE = re.compile(r"fila (\d+)")


def problems_report(sheet: Sheet, problems: list[tuple[str | None, str]]) -> bytes | None:
    """The ``.xlsx`` of the rows named in ``problems``; ``None`` if none is."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    reasons: dict[int, list[str]] = {}
    for source_id, message in problems:
        found = _LINE.fullmatch(source_id or "")
        if found:
            reasons.setdefault(int(found.group(1)), []).append(message)
    if not reasons:
        return None

    by_line = dict(zip(sheet.lines, sheet.rows, strict=True))
    book = Workbook()
    rows = book.active
    rows.title = "Filas con problemas"
    rows.append(["Fila", "Motivo", *sheet.headers])
    for cell in rows[1]:
        cell.font = Font(bold=True)
    highlight = PatternFill("solid", fgColor="FDE68A")
    for line in sorted(reasons):
        rows.append([line, " · ".join(reasons[line]), *by_line.get(line, [])])
        rows.cell(row=rows.max_row, column=2).fill = highlight
    rows.column_dimensions["A"].width = 8
    rows.column_dimensions["B"].width = 70
    notes = book.create_sheet("Instrucciones")
    for text in (
        "Estas son las filas de tu archivo que no entraron completas, con el motivo.",
        "Las que dicen «no importada» o «no se importa» no entraron; en las demás faltó solo el dato indicado.",
        "Corrige los datos en las columnas originales, borra las columnas Fila y Motivo, y sube este archivo.",
        "Lo que ya se importó no se duplica.",
    ):
        notes.append([text])
    notes.column_dimensions["A"].width = 110
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
