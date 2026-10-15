"""Read an uploaded spreadsheet — ``.xlsx`` or ``.csv`` — into rows of text.

The first non-empty row is the header. Every cell comes back as a string,
the way the person sees it, so what each value means is decided later, by
the field it is assigned to. Limits keep a stray file from filling the
server's memory; ``openpyxl`` parses with ``defusedxml`` when installed,
which it is, so an ``.xlsx`` built to attack the XML parser is refused.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 80
SUFFIXES = (".xlsx", ".csv")


class SheetError(ValueError):
    """The file cannot be read as a sheet of patients."""


@dataclass(frozen=True)
class Sheet:
    headers: list[str]
    #: Data rows, each as long as ``headers``.
    rows: list[list[str]]
    #: The line of the file each data row was on, for messages.
    lines: list[int]
    #: The first row looked like data, not names: columns are "Columna N"
    #: and that row is imported too.
    headerless: bool = False


def read_sheet(path: Path) -> Sheet:
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise SheetError("Solo se leen archivos .xlsx y .csv.")
    if path.stat().st_size > MAX_BYTES:
        raise SheetError(f"El archivo pesa más de {MAX_BYTES // (1024 * 1024)} MB.")
    raw = _read_xlsx(path) if suffix == ".xlsx" else _read_csv(path)
    return _tidy(raw)


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        # A date typed into Excel arrives as a datetime at midnight.
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # zip, XML or openpyxl's own errors
        raise SheetError(
            "No se pudo leer el .xlsx: puede estar dañado o protegido con contraseña."
        ) from exc
    try:
        rows: list[list[str]] = []
        for values in book.worksheets[0].iter_rows(values_only=True):
            rows.append([_cell(v) for v in values[:MAX_COLUMNS]])
            if len(rows) > MAX_ROWS + 1:
                raise SheetError(
                    f"La hoja tiene más de {MAX_ROWS} filas; divídela en varios archivos."
                )
        return rows
    finally:
        book.close()


def _read_csv(path: Path) -> list[list[str]]:
    from charset_normalizer import from_bytes

    data = path.read_bytes()
    text = _decode(data, from_bytes)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:20_000], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    for values in csv.reader(io.StringIO(text), dialect):
        rows.append([v.strip() for v in values[:MAX_COLUMNS]])
        if len(rows) > MAX_ROWS + 1:
            raise SheetError(
                f"El archivo tiene más de {MAX_ROWS} filas; divídelo en varios archivos."
            )
    return rows


#: What a CSV is saved in by the people who send them: Excel on Windows
#: (cp1252), Excel's "Unicode text" (UTF-16), old Macs and Latin-1. Left
#: free, detection on a short Spanish file picked cp1250 and wrote "Núñez"
#: as "Núńez"; a sheet in Spanish is never in a Central European page.
_LIKELY_ENCODINGS = ["cp1252", "utf_16", "latin_1", "mac_roman"]


def _decode(data: bytes, from_bytes: Any) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    best = from_bytes(data, cp_isolation=_LIKELY_ENCODINGS).best()
    text = str(best) if best is not None else data.decode("cp1252", errors="replace")
    return text.lstrip("\N{ZERO WIDTH NO-BREAK SPACE}")


def _tidy(raw: list[list[str]]) -> Sheet:
    numbered = [(line, row) for line, row in enumerate(raw, 1) if any(cell for cell in row)]
    if not numbered:
        raise SheetError("El archivo está vacío.")
    headerless = _looks_like_data(numbered[0][1])
    header, data = ([], numbered) if headerless else (numbered[0][1], numbered[1:])
    if not data:
        raise SheetError("El archivo solo tiene la fila de encabezados, sin datos debajo.")
    width = max(len(row) for _, row in numbered)
    headers: list[str] = []
    for i in range(width):
        name = header[i].strip() if i < len(header) else ""
        name = name or f"Columna {i + 1}"
        # Two columns with one name stay apart.
        while name in headers:
            name = f"{name} ({i + 1})"
        headers.append(name)
    return Sheet(
        headers=headers,
        rows=[row + [""] * (width - len(row)) for _, row in data],
        lines=[line for line, _ in data],
        headerless=headerless,
    )


_VALUE = re.compile(
    r"^(\S+@\S+\.\S+"  # an email
    r"|[+(]?\d[\d\s().\-/:,$]*"  # a number, phone, amount, date or time
    r"|[A-Z]{4}\d{6}[HM][A-Z]{5}[A-Z0-9]\d)$",  # a CURP
    re.IGNORECASE,
)


def _looks_like_data(row: list[str]) -> bool:
    """Whether a would-be header row is really the first patient: column
    names are words; emails, numbers, dates and ID documents are values."""
    cells = [cell for cell in row if cell]
    values = [cell for cell in cells if _VALUE.match(cell)]
    return len(values) >= 2 and len(values) * 2 >= len(cells)
